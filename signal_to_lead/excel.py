from __future__ import annotations

from pathlib import Path
from typing import Dict, Optional

from openpyxl import Workbook, load_workbook


EXCEL_HEADERS = [
    "lead_id",
    "email",
    "company",
    "country",
    "state",
    "employee_count",
    "revenue",
    "product_interest",
    "firmographic_score",
    "intent_score",
    "email_opens",
    "email_clicks",
    "site_visits",
    "last_activity_at",
    "score",
    "queue",
    "owner",
    "rule_match",
]


def append_row(path: str, row: Dict[str, object], sheet_name: str = "Leads") -> None:
    p = Path(path)
    if p.exists():
        wb = load_workbook(p)
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb.create_sheet(title=sheet_name)
            ws.append(EXCEL_HEADERS)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        ws.append(EXCEL_HEADERS)

    values = [
        row.get("lead_id"),
        row.get("email"),
        row.get("company"),
        row.get("country"),
        row.get("state"),
        row.get("employee_count"),
        row.get("revenue"),
        ",".join(row.get("product_interest", []) or []),
        row.get("firmographic_score"),
        row.get("intent_score"),
        row.get("email_opens"),
        row.get("email_clicks"),
        row.get("site_visits"),
        row.get("last_activity_at"),
        row.get("score"),
        row.get("queue"),
        row.get("owner"),
        row.get("rule_match"),
    ]
    ws.append(values)
    wb.save(p)
